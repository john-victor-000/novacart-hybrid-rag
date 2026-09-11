from pathlib import Path

import pandas as pd
import pymupdf
from docx import Document

from backend.app.ingestion import ingest_directory, ingest_file


def test_pdf_ingestion_extracts_pages(tmp_path: Path) -> None:
    pdf_path = tmp_path / "shipping_policy.pdf"
    pdf = pymupdf.open()
    first_page = pdf.new_page()
    first_page.insert_text((72, 72), "Shipping policy\nOrders ship in two days.")
    second_page = pdf.new_page()
    second_page.insert_text((72, 72), "Tracking details are emailed.")
    pdf.save(pdf_path)
    pdf.close()

    records = ingest_file(pdf_path)

    assert len(records) == 2
    assert records[0].document_name == "shipping_policy.pdf"
    assert records[0].document_type == "pdf"
    assert records[0].page == 1
    assert records[0].section is None
    assert "Orders ship in two days." in records[0].text
    assert records[1].page == 2
    assert "Tracking details" in records[1].text
    assert records[0].document_id != records[1].document_id


def test_docx_ingestion_extracts_sections(tmp_path: Path) -> None:
    docx_path = tmp_path / "warranty_guide.docx"
    doc = Document()
    doc.add_heading("Warranty Coverage", level=1)
    doc.add_paragraph("NovaCart covers manufacturing defects.")
    doc.add_heading("Exclusions", level=1)
    doc.add_paragraph("Accidental damage is excluded.")
    doc.save(docx_path)

    records = ingest_file(docx_path)

    assert len(records) == 2
    assert records[0].document_name == "warranty_guide.docx"
    assert records[0].document_type == "docx"
    assert records[0].page is None
    assert records[0].section == "Warranty Coverage"
    assert "manufacturing defects" in records[0].text
    assert records[1].section == "Exclusions"
    assert "Accidental damage" in records[1].text


def test_csv_ingestion_extracts_rows(tmp_path: Path) -> None:
    csv_path = tmp_path / "products.csv"
    pd.DataFrame(
        [
            {"sku": "NC-1", "name": "Nova Lamp", "description": "Desk lighting"},
            {"sku": "NC-2", "name": "Nova Chair", "description": "Ergonomic chair"},
        ]
    ).to_csv(csv_path, index=False)

    records = ingest_file(csv_path)

    assert len(records) == 2
    assert records[0].document_name == "products.csv"
    assert records[0].document_type == "csv"
    assert records[0].page is None
    assert records[0].section == "row 1"
    assert "sku: NC-1" in records[0].text
    assert "description: Desk lighting" in records[0].text
    assert records[1].section == "row 2"
    assert "Nova Chair" in records[1].text


def test_directory_ingestion_skips_unsupported_files(tmp_path: Path) -> None:
    csv_path = tmp_path / "products.csv"
    pd.DataFrame([{"sku": "NC-1", "name": "Nova Lamp"}]).to_csv(
        csv_path, index=False
    )
    (tmp_path / "notes.txt").write_text("not part of phase 2", encoding="utf-8")

    records = ingest_directory(tmp_path)

    assert len(records) == 1
    assert records[0].document_name == "products.csv"
