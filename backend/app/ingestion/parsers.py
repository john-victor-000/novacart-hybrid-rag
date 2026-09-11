"""Parsers for supported Phase 2 source document types."""

from collections.abc import Iterator
from hashlib import sha256
from pathlib import Path
import re

import pandas as pd
import pymupdf
from docx import Document

from backend.app.ingestion.models import IngestedDocument

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv"}


def parse_pdf(path: Path) -> list[IngestedDocument]:
    """Extract one normalized document per PDF page."""
    records: list[IngestedDocument] = []

    with pymupdf.open(path) as pdf:
        for page_index, page in enumerate(pdf, start=1):
            text = _normalize_text(page.get_text("text"))
            records.append(
                _build_record(
                    path=path,
                    document_type="pdf",
                    page=page_index,
                    section=None,
                    text=text,
                )
            )

    return records


def parse_docx(path: Path) -> list[IngestedDocument]:
    """Extract one normalized document per DOCX section."""
    doc = Document(path)
    sections: list[tuple[str | None, list[str]]] = []
    current_section: str | None = None
    current_parts: list[str] = []

    for block in _iter_docx_blocks(doc):
        style_name = getattr(block.style, "name", "") if hasattr(block, "style") else ""
        text = _normalize_text(block.text)
        if not text:
            continue

        if style_name.startswith("Heading"):
            if current_parts:
                sections.append((current_section, current_parts))
            current_section = text
            current_parts = []
            continue

        current_parts.append(text)

    if current_parts:
        sections.append((current_section, current_parts))

    if not sections and current_section:
        sections.append((current_section, [current_section]))

    return [
        _build_record(
            path=path,
            document_type="docx",
            page=None,
            section=section,
            text=_normalize_text("\n".join(parts)),
        )
        for section, parts in sections
    ]


def parse_csv(path: Path) -> list[IngestedDocument]:
    """Extract one normalized document per CSV row."""
    frame = pd.read_csv(path).fillna("")
    records: list[IngestedDocument] = []

    for row_index, row in frame.iterrows():
        row_number = int(row_index) + 1
        text = _normalize_text(
            "\n".join(f"{column}: {value}" for column, value in row.items())
        )
        records.append(
            _build_record(
                path=path,
                document_type="csv",
                page=None,
                section=f"row {row_number}",
                text=text,
            )
        )

    return records


def _iter_docx_blocks(doc: Document) -> Iterator[object]:
    for paragraph in doc.paragraphs:
        yield paragraph

    for table in doc.tables:
        for row in table.rows:
            cells = [_normalize_text(cell.text) for cell in row.cells]
            text = " | ".join(cell for cell in cells if cell)
            if text:
                yield _TableBlock(text=text)


class _TableBlock:
    def __init__(self, text: str) -> None:
        self.text = text
        self.style = None


def _build_record(
    path: Path,
    document_type: str,
    page: int | None,
    section: str | None,
    text: str,
) -> IngestedDocument:
    normalized_path = path.resolve()
    document_id = _document_id(normalized_path, page, section, text)

    return IngestedDocument(
        document_id=document_id,
        document_name=path.name,
        document_type=document_type,
        source=str(normalized_path),
        page=page,
        section=section,
        text=text,
    )


def _document_id(
    path: Path, page: int | None, section: str | None, text: str
) -> str:
    raw_id = f"{path}|{page or ''}|{section or ''}|{text}"
    return sha256(raw_id.encode("utf-8")).hexdigest()


def _normalize_text(value: object) -> str:
    return re.sub(r"[ \t]+", " ", str(value)).strip()
